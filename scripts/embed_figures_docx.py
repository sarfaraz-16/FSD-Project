import os
import docx
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

INPUT_DOCX = 'SmartERP_Project_Report.docx'
BACKUP_DOCX = 'SmartERP_Project_Report_Backup.docx'
OUTPUT_DOCX = 'SmartERP_Project_Report_Completed.docx'
IMG_DIR = os.path.join('docs', 'images')

FIGURE_MAPPING = {
    '[ Paste Figure 4.1 (Layered System Architecture) here ]': 'Figure_4_1_Layered_Architecture.png',
    '[ Paste Figure 4.2 (Role-Based Access Control Hierarchy) here ]': 'Figure_4_2_RBAC_Hierarchy.png',
    '[ Paste Figure 4.3 (UML Class Diagram of Document Models) here ]': 'Figure_4_3_UML_Class_Diagram.png',
    '[ Paste Figure 4.4 (Sales Order Confirmation Sequence Diagram) here ]': 'Figure_4_4_Sequence_Diagram.png',
    '[ Paste Figure 4.5 (Transactional Stock Update Algorithm Flowchart) here ]': 'Figure_4_5_Stock_Update_Flowchart.png',
    '[ Paste Figure A.1 (Login Page) here ]': 'Figure_A1_Login_Page.png',
    '[ Paste Figure A.2 (Executive Dashboard) here ]': 'Figure_A2_Executive_Dashboard.png',
    '[ Paste Figure A.3 (Inventory Product List) here ]': 'Figure_A3_Inventory_Products.png',
    '[ Paste Figure A.4 (Sales Order and Order Confirmation) here ]': 'Figure_A4_Sales_Order.png',
    '[ Paste Figure A.5 (Purchase Order and Receive Stock) here ]': 'Figure_A5_Purchase_Order.png',
    '[ Paste Figure A.6 (Reports and CSV Export) here ]': 'Figure_A6_Reports.png',
    '[ Paste Figure A.7 (Employee Directory (HR Module)) here ]': 'Figure_A7_Employee_Directory.png',
    '[ Paste Figure A.8 (Users and Roles Management) here ]': 'Figure_A8_Users_Management.png'
}

def process_paragraph(p, replaced_count):
    p_text = p.text.strip()
    for placeholder, img_filename in FIGURE_MAPPING.items():
        if placeholder in p_text:
            img_path = os.path.join(IMG_DIR, img_filename)
            if os.path.exists(img_path):
                p.text = ''
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(img_path, width=Inches(6.0))
                replaced_count[0] += 1
                print(f"Inserted: {img_filename} for placeholder: {placeholder}")
            else:
                print(f"Warning: Image file not found: {img_path}")

def main():
    if not os.path.exists(INPUT_DOCX):
        print(f"File not found: {INPUT_DOCX}")
        return

    doc = docx.Document(INPUT_DOCX)
    replaced_count = [0]

    # Check top-level paragraphs
    for p in doc.paragraphs:
        process_paragraph(p, replaced_count)

    # Check paragraphs inside tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    process_paragraph(p, replaced_count)

    # Backup original if not yet backed up
    if not os.path.exists(BACKUP_DOCX):
        import shutil
        shutil.copyfile(INPUT_DOCX, BACKUP_DOCX)
        print(f"Backup created at: {BACKUP_DOCX}")

    doc.save(OUTPUT_DOCX)
    print(f"Successfully processed and embedded {replaced_count[0]} figures into {OUTPUT_DOCX}!")

if __name__ == '__main__':
    main()
